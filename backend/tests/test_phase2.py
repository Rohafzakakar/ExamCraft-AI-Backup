import io
import pytest
import pytest_asyncio
import pymupdf
import docx
from httpx import ASGITransport, AsyncClient
from backend.app.core.database import Base, engine
from backend.app.main import app
from backend.app.services.chunker import DocumentChunker
from backend.app.services.embedding import get_embedding_service
from backend.app.services.parser import (
    DocumentParser,
    UnreadableDocumentError,
    UnsupportedFileTypeError,
)
from backend.app.services.rag_service import get_rag_service
from backend.app.services.vector_store import get_vector_repository


@pytest_asyncio.fixture(autouse=True)
async def prepare_db():
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    yield
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.drop_all)


def create_sample_pdf_bytes() -> bytes:
    """Creates an in-memory 2-page PDF document with chapters."""
    doc = pymupdf.open()
    # Page 1
    p1 = doc.new_page()
    p1.insert_text(
        pymupdf.Point(50, 72),
        "Chapter 1: Principles of Cellular Biology\n\n"
        "Photosynthesis is the fundamental biological process by which autotrophic organisms convert light energy into chemical energy stored in glucose molecules. "
        "The light-dependent reactions take place within the thylakoid membranes of chloroplasts, where chlorophyll pigments absorb photons. "
        "Water molecules undergo photolysis, yielding protons, electrons, and releasing molecular oxygen as a crucial metabolic byproduct.",
    )
    # Page 2
    p2 = doc.new_page()
    p2.insert_text(
        pymupdf.Point(50, 72),
        "Chapter 2: Genetics and Inheritance\n\n"
        "Gregor Mendel established the foundation of classical genetics through his hybridization experiments with Pisum sativum pea plants. "
        "The Law of Segregation posits that each diploid organism possesses two alleles for each gene, which separate during gamete formation. "
        "The Law of Independent Assortment explains the random distribution of alleles for unlinked genes.",
    )
    buf = io.BytesIO()
    doc.save(buf)
    doc.close()
    return buf.getvalue()


def create_sample_docx_bytes() -> bytes:
    """Creates an in-memory DOCX document with headings."""
    doc = docx.Document()
    doc.add_heading("Chapter 1: Operating Systems Concepts", level=1)
    doc.add_paragraph(
        "An operating system acts as an intermediary between computer hardware and user application programs. "
        "Process scheduling manages the allocation of CPU time across multiple active threads to maximize resource utilization and responsiveness."
    )
    doc.add_heading("Section 1.1: Memory Management", level=2)
    doc.add_paragraph(
        "Virtual memory decouples logical address spaces from physical RAM using paging and segmentation schemes. "
        "When a requested page is not present in physical frames, a page fault triggers retrieval from swap space on persistent storage."
    )
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_pdf_parsing_and_metadata():
    pdf_bytes = create_sample_pdf_bytes()
    parsed = DocumentParser.parse_pdf(pdf_bytes, "CellularBiology.pdf")

    assert parsed.total_pages == 2
    assert parsed.file_type == "pdf"
    assert parsed.detected_language == "english"
    assert parsed.is_rtl is False
    assert len(parsed.pages) == 2
    assert "Photosynthesis" in parsed.pages[0].text
    assert "Mendel" in parsed.pages[1].text
    assert parsed.pages[0].page_number == 1
    assert parsed.pages[1].page_number == 2


def test_docx_parsing_and_headings():
    docx_bytes = create_sample_docx_bytes()
    parsed = DocumentParser.parse_docx(docx_bytes, "OperatingSystems.docx")

    assert parsed.file_type == "docx"
    assert len(parsed.pages) >= 1
    assert parsed.detected_language == "english"
    assert "Operating Systems" in parsed.pages[0].chapter
    assert "Virtual memory" in parsed.pages[0].text


def test_unreadable_and_unsupported_files():
    # Empty file
    with pytest.raises(UnreadableDocumentError):
        DocumentParser.parse_pdf(b"", "empty.pdf")

    # Scanned / empty text file
    doc = pymupdf.open()
    doc.new_page()  # Blank page with no text
    buf = io.BytesIO()
    doc.save(buf)
    with pytest.raises(UnreadableDocumentError):
        DocumentParser.parse_pdf(buf.getvalue(), "blank.pdf")

    # Unsupported format
    with pytest.raises(UnsupportedFileTypeError):
        DocumentParser.parse_file(b"Sample plain text", "document.txt")


def test_multilingual_and_rtl_detection():
    # English
    lang, rtl = DocumentParser.detect_language_and_direction(
        "This is an academic syllabus for Computer Science at Kabul University."
    )
    assert lang == "english"
    assert rtl is False

    # Dari / Persian
    dari_text = "این یک سند آموزشی برای مضمون بیولوژی و ژنتیک مالیکولی در پوهنتون کابل است که ساختار حجره را تشریح میکند."
    lang, rtl = DocumentParser.detect_language_and_direction(dari_text)
    assert lang == "dari"
    assert rtl is True

    # Pashto
    pashto_text = "دا د بیولوژي تدریسي کتاب دی چې د رڼا ترکیب او د حجرو ویش په اړه مفصل معلومات وړاندې کوي او د ښوونځیو لپاره دی."
    lang, rtl = DocumentParser.detect_language_and_direction(pashto_text)
    assert lang in ["pashto", "arabic", "dari"]
    assert rtl is True


def test_chunking_with_strict_metadata_preservation():
    pdf_bytes = create_sample_pdf_bytes()
    parsed = DocumentParser.parse_pdf(pdf_bytes, "CellularBiology.pdf")

    chunker = DocumentChunker(target_tokens=100, overlap_tokens=20)
    all_chunks = chunker.chunk_document("doc_101", parsed)

    assert len(all_chunks) >= 2
    for c in all_chunks:
        assert c["document_id"] == "doc_101"
        assert c["page_number"] in [1, 2]
        assert c["token_count"] > 0
        assert len(c["content"]) > 10
        assert "chapter" in c

    # Test Source-Approved Selection (only Page 2 approved)
    page2_chunks = chunker.chunk_document("doc_101", parsed, selected_pages=[2])
    assert len(page2_chunks) >= 1
    for c in page2_chunks:
        assert c["page_number"] == 2
        assert "Mendel" in c["content"]
        assert "Photosynthesis" not in c["content"]


@pytest.mark.asyncio
async def test_multilingual_embedding_and_rag_retrieval():
    embedding_svc = get_embedding_service()
    vector_repo = get_vector_repository()
    rag_svc = get_rag_service()

    # Create and index chunks from our 2-page document
    pdf_bytes = create_sample_pdf_bytes()
    parsed = DocumentParser.parse_pdf(pdf_bytes, "CellularBiology.pdf")
    chunker = DocumentChunker()
    chunks = chunker.chunk_document("doc_bio", parsed)

    embeddings = embedding_svc.embed_batch([c["content"] for c in chunks])
    assert len(embeddings) == len(chunks)
    assert len(embeddings[0]) == 384

    await vector_repo.add_chunks(chunks, embeddings)

    # 1. Query for photosynthesis (Should match Page 1)
    retrieved = await rag_svc.retrieve_approved_context(
        query="light dependent thylakoid chloroplasts",
        document_id="doc_bio",
        document_name="CellularBiology.pdf",
        approved_pages=[1, 2],
        top_k=1,
    )
    assert len(retrieved) == 1
    assert retrieved[0].page_number == 1
    assert "Photosynthesis" in retrieved[0].content

    # 2. Query restricted ONLY to approved page 2 (Should never retrieve page 1 even if mentioned)
    retrieved_page2 = await rag_svc.retrieve_approved_context(
        query="cellular processes and biology",
        document_id="doc_bio",
        document_name="CellularBiology.pdf",
        approved_pages=[2],
        top_k=1,
    )
    assert len(retrieved_page2) == 1
    assert retrieved_page2[0].page_number == 2

    # 3. Prompt context assembly includes citations
    prompt_ctx = rag_svc.assemble_prompt_context(retrieved)
    assert "SOURCE EXCERPT" in prompt_ctx
    assert "Page: 1" in prompt_ctx


@pytest.mark.asyncio
async def test_api_document_upload_and_chunk_endpoints():
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        pdf_bytes = create_sample_pdf_bytes()

        # Upload document
        files = {"file": ("BiologyExamMaterial.pdf", pdf_bytes, "application/pdf")}
        data = {"session_id": "guest_test_session"}
        res = await client.post("/api/documents/upload", files=files, data=data)
        assert res.status_code == 201
        doc_data = res.json()
        doc_id = doc_data["id"]

        assert doc_data["filename"] == "BiologyExamMaterial.pdf"
        assert doc_data["total_pages"] == 2
        assert doc_data["chunks_count"] >= 2
        assert doc_data["status"] == "indexed"

        # List documents for guest session
        list_res = await client.get("/api/documents?session_id=guest_test_session")
        assert list_res.status_code == 200
        assert len(list_res.json()) >= 1

        # Fetch document chunks
        chunks_res = await client.get(f"/api/documents/{doc_id}/chunks")
        assert chunks_res.status_code == 200
        chunks = chunks_res.json()
        assert len(chunks) >= 2

        # Fetch chunks filtered by approved page 1
        page1_chunks_res = await client.get(f"/api/documents/{doc_id}/chunks?pages=1")
        assert page1_chunks_res.status_code == 200
        p1_chunks = page1_chunks_res.json()
        assert len(p1_chunks) >= 1
        for c in p1_chunks:
            assert c["page_number"] == 1
