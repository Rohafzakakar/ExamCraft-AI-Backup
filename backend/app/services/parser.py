import io
import os
import re
import unicodedata
from dataclasses import dataclass, field
from typing import Any, Dict, List, Optional, Tuple
import pymupdf
import docx
from backend.app.core.config import settings


class DocumentProcessingError(Exception):
    """Base exception for document parsing failures."""
    pass


class UnsupportedFileTypeError(DocumentProcessingError):
    pass


class FileTooLargeError(DocumentProcessingError):
    pass


class CorruptedFileError(DocumentProcessingError):
    pass


class UnreadableDocumentError(DocumentProcessingError):
    pass


@dataclass
class DocumentPage:
    page_number: int
    text: str
    chapter: Optional[str] = None
    section: Optional[str] = None
    char_count: int = 0
    tables: List[List[List[str]]] = field(default_factory=list)


@dataclass
class ParsedDocument:
    filename: str
    file_type: str  # 'pdf' or 'docx'
    file_size_bytes: int
    total_pages: int
    detected_language: str
    is_rtl: bool
    pages: List[DocumentPage]
    chapters: List[str]
    metadata: Dict[str, Any] = field(default_factory=dict)
    total_characters: int = 0


class DocumentParser:
    """Parses PDF and DOCX educational documents with strict metadata preservation."""

    @staticmethod
    def detect_language_and_direction(text: str) -> Tuple[str, bool]:
        """Detects the primary language and text direction (RTL vs LTR) of the document."""
        if not text or len(text.strip()) == 0:
            return "auto", False

        arabic_script_chars = 0
        latin_chars = 0
        cyrillic_chars = 0
        total_letters = 0

        pashto_markers = set("ښځڅډړږڼۍئ")
        dari_persian_markers = set("گچپژ")
        urdu_markers = set("ٹڈڑںےہۂ")

        has_pashto = False
        has_dari = False
        has_urdu = False

        for ch in text[:50000]:
            cp = ord(ch)
            if unicodedata.category(ch).startswith("L"):
                total_letters += 1
                if 0x0600 <= cp <= 0x06FF or 0x0750 <= cp <= 0x077F or 0x08A0 <= cp <= 0x08FF:
                    arabic_script_chars += 1
                    if ch in pashto_markers:
                        has_pashto = True
                    if ch in dari_persian_markers:
                        has_dari = True
                    if ch in urdu_markers:
                        has_urdu = True
                elif 0x0041 <= cp <= 0x005A or 0x0061 <= cp <= 0x007A or 0x00C0 <= cp <= 0x024F:
                    latin_chars += 1
                elif 0x0400 <= cp <= 0x04FF:
                    cyrillic_chars += 1

        if total_letters == 0:
            return "auto", False

        rtl_ratio = arabic_script_chars / total_letters

        if rtl_ratio > 0.25:
            if has_pashto:
                return "pashto", True
            if has_urdu:
                return "urdu", True
            if has_dari:
                return "dari", True
            return "arabic", True

        if latin_chars / total_letters > 0.5:
            return "english", False

        return "auto", False

    @classmethod
    def parse_pdf(cls, file_bytes: bytes, filename: str) -> ParsedDocument:
        """Extracts text and hierarchical chapter/page structure from a PDF file."""
        if len(file_bytes) == 0:
            raise UnreadableDocumentError("The uploaded PDF file is empty (0 bytes).")
        if len(file_bytes) > settings.UPLOAD_MAX_BYTES:
            raise FileTooLargeError(f"File size exceeds maximum allowed limit of {settings.UPLOAD_MAX_BYTES // (1024 * 1024)}MB.")

        try:
            doc = pymupdf.open(stream=file_bytes, filetype="pdf")
        except Exception as e:
            raise CorruptedFileError(f"Could not open PDF file. The file may be damaged or password protected: {str(e)}")

        pages: List[DocumentPage] = []
        chapters_set: List[str] = []
        all_text_list: List[str] = []

        toc = doc.get_toc()
        page_to_chapter: Dict[int, str] = {}
        if toc:
            current_ch = None
            for item in toc:
                if len(item) >= 3:
                    lvl, title, pno = item[0], item[1], item[2]
                    if lvl == 1:
                        current_ch = title
                        if title not in chapters_set:
                            chapters_set.append(title)
                    page_to_chapter[pno] = current_ch or title

        current_chapter = "Introduction"
        current_section = None
        total_pages = doc.page_count
        if total_pages == 0:
            raise UnreadableDocumentError("The uploaded PDF contains no pages.")

        for pno in range(total_pages):
            page_idx = pno + 1
            page = doc.load_page(pno)
            text = page.get_text("text") or ""
            text = text.strip()

            if page_idx in page_to_chapter:
                current_chapter = page_to_chapter[page_idx]

            lines = [line.strip() for line in text.split("\n") if line.strip()]
            for line in lines[:5]:
                ch_match = re.match(r"^(Chapter\s+[0-9IVXLCDM]+|فصل\s+[۰-۹0-9]+|باب\s+[۰-۹0-9]+)[:\s\.\-]*(.*)", line, re.IGNORECASE)
                if ch_match:
                    current_chapter = line
                    if line not in chapters_set:
                        chapters_set.append(line)
                    break
                sec_match = re.match(r"^([0-9]+\.[0-9]+|Section\s+[0-9]+)[:\s\.\-]*(.*)", line, re.IGNORECASE)
                if sec_match:
                    current_section = line

            pages.append(
                DocumentPage(
                    page_number=page_idx,
                    text=text,
                    chapter=current_chapter,
                    section=current_section,
                    char_count=len(text),
                )
            )
            if text:
                all_text_list.append(text)

        doc.close()

        full_text = "\n\n".join(all_text_list)
        if len(full_text.strip()) < 50:
            raise UnreadableDocumentError(
                "The selected document does not contain readable text. Scanned images or empty documents cannot be processed."
            )

        lang, is_rtl = cls.detect_language_and_direction(full_text)
        if not chapters_set and current_chapter:
            chapters_set = [current_chapter]

        return ParsedDocument(
            filename=filename,
            file_type="pdf",
            file_size_bytes=len(file_bytes),
            total_pages=total_pages,
            detected_language=lang,
            is_rtl=is_rtl,
            pages=pages,
            chapters=chapters_set or ["General Content"],
            metadata={"toc_found": bool(toc), "page_count": total_pages},
            total_characters=len(full_text),
        )

    @classmethod
    def parse_docx(cls, file_bytes: bytes, filename: str) -> ParsedDocument:
        """Extracts text, headings, and structure from a DOCX file."""
        if len(file_bytes) == 0:
            raise UnreadableDocumentError("The uploaded DOCX file is empty (0 bytes).")
        if len(file_bytes) > settings.UPLOAD_MAX_BYTES:
            raise FileTooLargeError(f"File size exceeds maximum allowed limit of {settings.UPLOAD_MAX_BYTES // (1024 * 1024)}MB.")

        try:
            doc = docx.Document(io.BytesIO(file_bytes))
        except Exception as e:
            raise CorruptedFileError(f"Could not open DOCX file: {str(e)}")

        pages: List[DocumentPage] = []
        chapters_set: List[str] = []
        all_text_list: List[str] = []

        current_chapter = "Chapter 1"
        current_section = None
        current_page_num = 1
        current_page_text: List[str] = []
        words_on_page = 0
        approx_words_per_page = 350

        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue

            style_name = para.style.name.lower() if para.style else ""

            if "heading 1" in style_name or "title" in style_name:
                current_chapter = text
                if text not in chapters_set:
                    chapters_set.append(text)
            elif "heading 2" in style_name or "heading 3" in style_name:
                current_section = text

            words = len(text.split())
            words_on_page += words
            current_page_text.append(text)

            if words_on_page >= approx_words_per_page:
                p_text = "\n\n".join(current_page_text)
                pages.append(
                    DocumentPage(
                        page_number=current_page_num,
                        text=p_text,
                        chapter=current_chapter,
                        section=current_section,
                        char_count=len(p_text),
                    )
                )
                all_text_list.append(p_text)
                current_page_num += 1
                current_page_text = []
                words_on_page = 0

        if current_page_text:
            p_text = "\n\n".join(current_page_text)
            pages.append(
                DocumentPage(
                    page_number=current_page_num,
                    text=p_text,
                    chapter=current_chapter,
                    section=current_section,
                    char_count=len(p_text),
                )
            )
            all_text_list.append(p_text)

        full_text = "\n\n".join(all_text_list)
        if len(full_text.strip()) < 50:
            raise UnreadableDocumentError(
                "The selected document does not contain readable text. Scanned images or empty documents cannot be processed."
            )

        lang, is_rtl = cls.detect_language_and_direction(full_text)

        return ParsedDocument(
            filename=filename,
            file_type="docx",
            file_size_bytes=len(file_bytes),
            total_pages=len(pages),
            detected_language=lang,
            is_rtl=is_rtl,
            pages=pages,
            chapters=chapters_set or ["General Content"],
            metadata={"paragraph_count": len(doc.paragraphs)},
            total_characters=len(full_text),
        )

    @classmethod
    def parse_file(cls, file_bytes: bytes, filename: str) -> ParsedDocument:
        """Auto-detects file type by extension and parses the document."""
        lower_name = filename.lower()
        if lower_name.endswith(".pdf"):
            return cls.parse_pdf(file_bytes, filename)
        elif lower_name.endswith(".docx"):
            return cls.parse_docx(file_bytes, filename)
        else:
            raise UnsupportedFileTypeError(
                f"Unsupported file format: '{os.path.splitext(filename)[1]}'. ExamCraft AI supports PDF (.pdf) and DOCX (.docx) documents."
            )
